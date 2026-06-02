import os
import subprocess
import pandas as pd
from langchain_community.document_loaders import (
    PyMuPDFLoader,
    TextLoader,
    Docx2txtLoader,
    UnstructuredWordDocumentLoader,
)
from langchain_core.documents import Document as LCDocument
from pptx import Presentation


def _load_pptx(filepath: str) -> list:
    """Load PPTX using python-pptx (no unstructured dependency)."""
    prs = Presentation(filepath)
    filename = os.path.basename(filepath)
    docs = []
    full_text = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        slide_texts.append(t)
        if slide_texts:
            text = f"Slide {slide_num}:\n" + "\n".join(slide_texts)
            docs.append(LCDocument(page_content=text, metadata={"filename": filename, "slide": slide_num}))
            full_text.append(text)
    if not docs:
        docs.append(LCDocument(page_content="(no text content)", metadata={"filename": filename}))
    return docs


EXTENSION_LOADERS = {
    ".pdf": PyMuPDFLoader,
    ".txt": TextLoader,
    ".doc": UnstructuredWordDocumentLoader,
    ".docx": Docx2txtLoader,
    ".md": TextLoader,
    ".ppt": None,   # legacy .ppt not supported — will raise
    ".pptx": None,  # handled via _load_pptx
}

EXCEL_EXTENSIONS = {".xlsx", ".xls", ".xlsm", ".csv"}


def _load_excel(filepath: str) -> list:
    """Load Excel/CSV file and return LangChain Document objects (one per sheet)."""
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".csv":
        df = pd.read_csv(filepath)
        filename = os.path.basename(filepath)
        text = df.to_csv(sep="\t", index=False, header=True)
        return [LCDocument(page_content=text, metadata={"filename": filename})]

    engine = "xlrd" if ext == ".xls" else "openpyxl"
    xls = pd.ExcelFile(filepath, engine=engine)
    docs = []
    filename = os.path.basename(filepath)
    for sheet_name in xls.sheet_names:
        df = xls.parse(sheet_name)
        lines = [df.to_csv(sep="\t", index=False, header=True)]
        text = f"Sheet: {sheet_name}\n" + "\n".join(lines)
        docs.append(LCDocument(page_content=text, metadata={"filename": filename, "sheet": sheet_name}))
    xls.close()
    return docs


def load_single_document(filepath: str) -> list:
    """Load a single document file and return LangChain Document objects."""
    ext = os.path.splitext(filepath)[1].lower()

    if ext in EXCEL_EXTENSIONS:
        return _load_excel(filepath)

    # Handle PPTX via python-pptx (no unstructured dependency)
    if ext == ".pptx":
        return _load_pptx(filepath)

    if ext == ".ppt":
        raise ValueError("Legacy .ppt format is not supported. Please convert to .pptx.")

    loader_cls = EXTENSION_LOADERS.get(ext)
    if loader_cls is None:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext in (".txt", ".md"):
        loader = loader_cls(filepath, encoding="utf-8")
    elif ext == ".doc":
        loader = loader_cls(filepath, mode="single")
    else:
        loader = loader_cls(filepath)

    docs = loader.load()
    filename = os.path.basename(filepath)
    for doc in docs:
        doc.metadata["filename"] = filename
    return docs


def get_content_preview(docs: list, max_chars: int = 500) -> str:
    text = " ".join(doc.page_content[:max_chars] for doc in docs)
    text = text.strip()[:max_chars]
    if not text:
        return "（此文档可能为图片扫描件，无法提取文本内容，请下载查看）"
    return text

def load_document_content(filepath: str) -> str:
    """Load full text content of a document for preview."""
    docs = load_single_document(filepath)
    return "\n\n".join(doc.page_content for doc in docs)


def convert_to_pdf(filepath: str) -> str | None:
    """Convert PPTX/DOCX to PDF via LibreOffice headless. Returns PDF path or None."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext not in (".pptx", ".ppt", ".docx", ".doc"):
        return None
    out_dir = os.path.join(os.path.dirname(filepath), ".pdf_cache")
    os.makedirs(out_dir, exist_ok=True)
    pdf_name = os.path.splitext(os.path.basename(filepath))[0] + ".pdf"
    pdf_path = os.path.join(out_dir, pdf_name)
    if os.path.exists(pdf_path):
        return pdf_path
    try:
        soffice = os.environ.get("SOFFICE_PATH", "soffice")
        if os.name == "nt" and soffice == "soffice":
            for base in [os.environ.get("ProgramFiles", "C:/Program Files"), os.environ.get("ProgramFiles(x86)", "C:/Program Files (x86)")]:
                path = os.path.join(base, "LibreOffice/program/soffice.exe")
                if os.path.exists(path):
                    soffice = path
                    break
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, filepath],
            check=True, capture_output=True, timeout=120,
            env={**os.environ, "HOME": os.path.expanduser("~")},
        )
        if os.path.exists(pdf_path):
            return pdf_path
    except Exception:
        pass
    return None
